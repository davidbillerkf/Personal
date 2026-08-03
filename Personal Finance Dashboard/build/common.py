"""Shared styling constants and helpers for the Finance Dashboard workbook build."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule, DataBarRule, IconSetRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker

FONT_NAME = "Arial"

# Palette
NAVY = "0B1F3A"
NAVY_DARK = "081428"
WHITE = "FFFFFF"
LIGHT_GRAY = "F3F4F6"
MED_GRAY = "E5E7EB"
DARK_TEXT = "1F2937"
GREEN = "15803D"
GREEN_LIGHT = "DCFCE7"
RED = "B91C1C"
RED_LIGHT = "FEE2E2"
ORANGE = "C2650B"
ORANGE_LIGHT = "FEF3C7"
BLUE_ACCENT = "2563EB"
BLUE_LIGHT = "DBEAFE"
PURPLE = "6D28D9"

def f(size=10, bold=False, color=DARK_TEXT, italic=False, name=FONT_NAME):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)

def fill(color):
    return PatternFill(fill_type="solid", start_color=color, end_color=color)

THIN = Side(style="thin", color="D1D5DB")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=Side(style="thin", color="D1D5DB"))

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)

CUR_FMT = '$#,##0.00;[RED]($#,##0.00)'
CUR_FMT0 = '$#,##0;[RED]($#,##0)'
PCT_FMT = '0.0%'
DATE_FMT = 'mm/dd/yyyy'


def style_sheet_title(ws, title, subtitle=None, span_cols=12):
    """Dark navy banner title bar at top of a sheet, rows 1-2."""
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=span_cols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = f(20, bold=True, color=WHITE)
    c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, span_cols + 1):
        ws.cell(row=1, column=col).fill = fill(NAVY)
        ws.cell(row=2, column=col).fill = fill(NAVY)
    if subtitle:
        c.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        c2 = ws.cell(row=1, column=1)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 10


def style_section_header(ws, row, col, text, span=6, size=12):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    c = ws.cell(row=row, column=col, value=text)
    c.font = f(size, bold=True, color=WHITE)
    c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 20
    for cc in range(col, col + span):
        ws.cell(row=row, column=cc).fill = fill(NAVY)
    return row + 1


def add_table(ws, start_row, start_col, headers, name, data_rows=None, style="TableStyleMedium9",
              col_widths=None, number_formats=None):
    """Write a header row + optional data rows, wrap as an Excel Table, return dict of info."""
    data_rows = data_rows or []
    n_cols = len(headers)
    for j, h in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + j, value=h)
        cell.font = f(10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = CENTER_WRAP
    for i, row_vals in enumerate(data_rows):
        for j, val in enumerate(row_vals):
            cell = ws.cell(row=start_row + 1 + i, column=start_col + j, value=val)
            if number_formats and j in number_formats:
                cell.number_format = number_formats[j]
            cell.font = f(10)
    end_row = start_row + max(len(data_rows), 1)
    end_col = start_col + n_cols - 1
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showFirstColumn=False)
    ws.add_table(tab)
    if col_widths:
        for j, w in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(start_col + j)].width = w
    return {"name": name, "start_row": start_row, "end_row": end_row, "start_col": start_col,
            "end_col": end_col, "headers": headers, "ref": ref, "first_data_row": start_row + 1}


def kpi_card(ws, row, col, label, formula, number_format=CUR_FMT0, width_cols=2, sub_label=None,
             value_color=DARK_TEXT):
    """A white KPI card spanning width_cols columns, 4 rows tall, with light border."""
    r0, c0 = row, col
    r1, c1 = row + 3, col + width_cols - 1
    for rr in range(r0, r1 + 1):
        for cc in range(c0, c1 + 1):
            cell = ws.cell(row=rr, column=cc)
            cell.fill = fill(WHITE)
            cell.border = BORDER_ALL
    ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c1)
    lab = ws.cell(row=r0, column=c0, value=label.upper())
    lab.font = f(9, bold=True, color="6B7280")
    lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 2, end_column=c1)
    val = ws.cell(row=r0 + 1, column=c0, value=formula)
    val.font = f(18, bold=True, color=value_color)
    val.number_format = number_format
    val.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if sub_label:
        ws.merge_cells(start_row=r0 + 3, start_column=c0, end_row=r0 + 3, end_column=c1)
        s = ws.cell(row=r0 + 3, column=c0, value=sub_label)
        s.font = f(8, italic=True, color="9CA3AF")
        s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return {"value_cell": val.coordinate}


def add_list_validation(ws, cell_range, formula1, allow_blank=True, error_title="Invalid Entry",
                         error_msg="Please choose a value from the dropdown list."):
    # Data validation formulas must not carry a leading "=" in the underlying XML (same rule as
    # cell <f> elements) — a leading "=" makes Excel silently fail to parse it, so no dropdown
    # arrow ever appears. Strip it here so every caller can still write "=Name" naturally.
    if formula1.startswith("="):
        formula1 = formula1[1:]
    dv = DataValidation(type="list", formula1=formula1, allow_blank=allow_blank,
                         showErrorMessage=True, errorTitle=error_title, error=error_msg)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv
