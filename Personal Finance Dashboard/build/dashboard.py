"""ChartData helper sheet + Dashboard sheet (simplified planner scope)."""
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from common import *
import seed_data as sd

EXPENSE_CATEGORIES = [c[1] for c in sd.CATEGORIES if c[2] == "Expense"]


def build_chartdata(wb):
    ws = wb.create_sheet("ChartData")
    headers = ["Month", "MonthEnd", "Income", "Expenses", "NetSavings", "CumSavings"]
    for j, h in enumerate(headers):
        ws.cell(row=1, column=1 + j, value=h).font = f(10, bold=True, color=WHITE)
        ws.cell(row=1, column=1 + j).fill = fill(NAVY)
    n = 13
    for i in range(n):
        months_ago = n - 1 - i
        rr = 2 + i
        ws.cell(row=rr, column=1, value=f"=EDATE(TodayDate,{-months_ago})")
        ws.cell(row=rr, column=1).number_format = "mmm-yy"
        ws.cell(row=rr, column=2, value=f"=EOMONTH(A{rr},0)")
        ws.cell(row=rr, column=2).number_format = DATE_FMT
        start_expr = f"EOMONTH(A{rr},-1)+1"
        ws.cell(row=rr, column=3,
                value=(f'=SUMIFS(Fact_Income[NetAmount],Fact_Income[Date],">="&{start_expr},Fact_Income[Date],"<="&B{rr})'
                       f'+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Income",BankImportRaw[Date],">="&{start_expr},BankImportRaw[Date],"<="&B{rr})'))
        ws.cell(row=rr, column=4,
                value=(f'=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[Date],">="&{start_expr},Fact_Expenses[Date],"<="&B{rr})'
                       f'+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Expense",BankImportRaw[Date],">="&{start_expr},BankImportRaw[Date],"<="&B{rr})'
                       f'-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Return",BankImportRaw[Date],">="&{start_expr},BankImportRaw[Date],"<="&B{rr})'))
        ws.cell(row=rr, column=5, value=f"=C{rr}-D{rr}")
        ws.cell(row=rr, column=6, value=f"=SUM($E$2:E{rr})")
        for c in (3, 4, 5, 6):
            ws.cell(row=rr, column=c).number_format = CUR_FMT0
    ws.sheet_state = "hidden"
    return {"first_row": 2, "last_row": 1 + n, "n": n}


def build_dashboard(wb, ctx):
    ws = wb.create_sheet("Dashboard", 0)
    style_sheet_title(ws, "Personal Finance Planner",
                       subtitle="Live financial command center — refresh data, review insights", span_cols=16)
    ws.sheet_view.showGridLines = False
    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 11.5

    ws.cell(row=3, column=1,
            value='="Data as of " & TEXT(TodayDate,"mmmm d, yyyy") & "  |  Refresh: Data > Refresh All, then Ctrl+Alt+F9 to recalculate"')
    ws.cell(row=3, column=1).font = f(9, italic=True, color="6B7280")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=10)

    # ---------------- KPI CARDS (8) ----------------
    kpis = [
        ("Monthly Income",
         '=SUMIFS(Fact_Income[NetAmount],Fact_Income[Date],">="&CurMonthStart,Fact_Income[Date],"<="&CurMonthEnd)'
         '+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Income",BankImportRaw[Date],">="&CurMonthStart,BankImportRaw[Date],"<="&CurMonthEnd)',
         CUR_FMT0, GREEN),
        ("Monthly Expenses",
         '=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[Date],">="&CurMonthStart,Fact_Expenses[Date],"<="&CurMonthEnd)'
         '+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Expense",BankImportRaw[Date],">="&CurMonthStart,BankImportRaw[Date],"<="&CurMonthEnd)'
         '-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Return",BankImportRaw[Date],">="&CurMonthStart,BankImportRaw[Date],"<="&CurMonthEnd)',
         CUR_FMT0, RED),
        ("Monthly Savings", "=A6-E6", CUR_FMT0, DARK_TEXT),
        ("Savings Rate %", "=IFERROR(I6/A6,0)", PCT_FMT, DARK_TEXT),
        ("Net Cash Flow", '=I6-SUMIFS(Fact_SavingsTransfers[Amount],Fact_SavingsTransfers[Date],">="&CurMonthStart,Fact_SavingsTransfers[Date],"<="&CurMonthEnd)', CUR_FMT0, DARK_TEXT),
        ("Budget Utilization %", "=IFERROR(SUM(Budget[Actual])/SUM(Budget[MonthlyBudget]),0)", PCT_FMT, ORANGE),
        ("Emergency Fund", '=SUMIFS(Fact_SavingsTransfers[Amount],Fact_SavingsTransfers[GoalName],"Emergency Fund")', CUR_FMT0, GREEN),
        ("Total Saved (All Goals)", "=SUM(Dim_Goal[CurrentAmount])", CUR_FMT0, GREEN),
    ]
    start_row = 5
    for i, (label, formula, fmt, color) in enumerate(kpis):
        row = start_row + (i // 4) * 5
        col = 1 + (i % 4) * 4
        kpi_card(ws, row, col, label, formula, number_format=fmt, width_cols=4, value_color=color)

    charts_top = start_row + 2 * 5 + 1  # after 2 rows of KPI cards

    def rng(sheet, col, r1, r2):
        return Reference(wb[sheet], min_col=col, min_row=r1, max_col=col, max_row=r2)

    cd = ctx["chartdata"]
    cd_first, cd_last = cd["first_row"], cd["last_row"]

    # ---- Chart 1: Income vs Expenses (13-month) ----
    ch1 = BarChart()
    ch1.type = "col"
    ch1.title = "Income vs Expenses — Trailing 13 Months"
    ch1.y_axis.title = "USD"
    ch1.height, ch1.width = 8, 16
    cats = rng("ChartData", 1, cd_first, cd_last)
    ch1.add_data(rng("ChartData", 3, cd_first - 1, cd_last), titles_from_data=True)
    ch1.add_data(rng("ChartData", 4, cd_first - 1, cd_last), titles_from_data=True)
    ch1.set_categories(cats)
    ch1.series[0].graphicalProperties.solidFill = GREEN
    ch1.series[1].graphicalProperties.solidFill = RED
    ws.add_chart(ch1, f"A{charts_top}")

    # ---- Chart 2: Spending Breakdown (pie, this month, from Budget) ----
    ch2 = PieChart()
    ch2.title = "Spending Breakdown — This Month by Category"
    ch2.height, ch2.width = 8, 16
    b_info = ctx["bud_info"]
    bfr, blr = b_info["first_data_row"], b_info["end_row"]
    data2 = rng("Budget", 3, bfr - 1, blr)
    cats2 = rng("Budget", 1, bfr, blr)
    ch2.add_data(data2, titles_from_data=True)
    ch2.set_categories(cats2)
    ws.add_chart(ch2, f"I{charts_top}")

    charts_row2 = charts_top + 17
    # ---- Chart 3: Spending Trend (line) ----
    ch3 = LineChart()
    ch3.title = "Spending Trend — Trailing 13 Months"
    ch3.height, ch3.width = 8, 16
    ch3.add_data(rng("ChartData", 4, cd_first - 1, cd_last), titles_from_data=True)
    ch3.set_categories(cats)
    ch3.series[0].graphicalProperties.line.solidFill = RED
    ch3.series[0].marker.symbol = "circle"
    ws.add_chart(ch3, f"A{charts_row2}")

    # ---- Chart 4: Savings Growth (cumulative) ----
    ch4 = LineChart()
    ch4.title = "Savings Growth — Cumulative"
    ch4.height, ch4.width = 8, 16
    ch4.add_data(rng("ChartData", 6, cd_first - 1, cd_last), titles_from_data=True)
    ch4.set_categories(cats)
    ch4.series[0].graphicalProperties.line.solidFill = GREEN
    ws.add_chart(ch4, f"I{charts_row2}")

    charts_row3 = charts_row2 + 17
    # ---- Chart 5: Budget Performance (actual vs budget) ----
    ch5 = BarChart()
    ch5.type = "col"
    ch5.title = "Budget Performance — Actual vs Budget"
    ch5.height, ch5.width = 8, 16
    ch5.add_data(rng("Budget", 2, bfr - 1, blr), titles_from_data=True)
    ch5.add_data(rng("Budget", 3, bfr - 1, blr), titles_from_data=True)
    ch5.set_categories(cats2)
    ch5.series[0].graphicalProperties.solidFill = "9CA3AF"
    ch5.series[1].graphicalProperties.solidFill = BLUE_ACCENT
    ws.add_chart(ch5, f"A{charts_row3}")

    # ---------------- WIDGETS ----------------
    widget_row = charts_row3 + 17
    widget_row = style_section_header(ws, widget_row, 1, "Widgets & Quick Insights", span=16, size=13)

    wr = widget_row + 1
    # -- Top Expenses (this month) --
    ws.cell(row=wr, column=1, value="Top Expenses (This Month, Expenses tab only)").font = f(10, bold=True)
    for j, h in enumerate(["Vendor", "Category", "Amount"]):
        ws.cell(row=wr + 1, column=1 + j, value=h).font = f(9, bold=True, color=WHITE)
        ws.cell(row=wr + 1, column=1 + j).fill = fill(NAVY)
    for i in range(5):
        rr = wr + 2 + i
        amt_cell = ws.cell(row=rr, column=3, value=f'=IFERROR(LARGE(Fact_Expenses[CurMonthAmt],{i+1}),0)')
        amt_cell.number_format = CUR_FMT
        ws.cell(row=rr, column=1,
                value=f'=IFERROR(INDEX(Fact_Expenses[VendorName],MATCH($C{rr},Fact_Expenses[CurMonthAmt],0)),"")')
        ws.cell(row=rr, column=2,
                value=f'=IFERROR(INDEX(Fact_Expenses[CategoryName],MATCH($C{rr},Fact_Expenses[CurMonthAmt],0)),"")')
        for c in (1, 2, 3):
            ws.cell(row=rr, column=c).font = f(9)

    # -- Largest Vendors (top 5, this month) --
    ws.cell(row=wr, column=5, value="Largest Vendors (This Month)").font = f(10, bold=True)
    for j, h in enumerate(["Vendor", "Amount"]):
        ws.cell(row=wr + 1, column=5 + j, value=h).font = f(9, bold=True, color=WHITE)
        ws.cell(row=wr + 1, column=5 + j).fill = fill(NAVY)
    for i in range(5):
        rr = wr + 2 + i
        amt_cell = ws.cell(row=rr, column=6, value=f'=IFERROR(LARGE(Dim_Vendor[CurMonthSpend],{i+1}),0)')
        amt_cell.number_format = CUR_FMT
        ws.cell(row=rr, column=5,
                value=f'=IFERROR(INDEX(Dim_Vendor[VendorName],MATCH($F{rr},Dim_Vendor[CurMonthSpend],0)),"")')
        for c in (5, 6):
            ws.cell(row=rr, column=c).font = f(9)

    # -- Largest Categories (top 5, this month, from Budget!Actual) --
    ws.cell(row=wr, column=8, value="Largest Categories (This Month)").font = f(10, bold=True)
    for j, h in enumerate(["Category", "Amount"]):
        ws.cell(row=wr + 1, column=8 + j, value=h).font = f(9, bold=True, color=WHITE)
        ws.cell(row=wr + 1, column=8 + j).fill = fill(NAVY)
    for i in range(5):
        rr = wr + 2 + i
        amt_cell = ws.cell(row=rr, column=9, value=f'=IFERROR(LARGE(Budget[Actual],{i+1}),0)')
        amt_cell.number_format = CUR_FMT
        ws.cell(row=rr, column=8,
                value=f'=IFERROR(INDEX(Budget[CategoryName],MATCH($I{rr},Budget[Actual],0)),"")')
        for c in (8, 9):
            ws.cell(row=rr, column=c).font = f(9)

    # -- Behavior Metrics --
    ws.cell(row=wr, column=11, value="Behavior Metrics").font = f(10, bold=True)
    metrics = [
        ("Need vs Want %",
         '=IFERROR((SUMIFS(Fact_Expenses[Amount],Fact_Expenses[NeedWantFlag],"Need")'
         '+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[NeedWantFlag],"Need",BankImportRaw[TransactionType],"Expense")'
         '-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[NeedWantFlag],"Need",BankImportRaw[TransactionType],"Return"))'
         '/(SUM(Fact_Expenses[Amount])'
         '+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Expense")'
         '-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[TransactionType],"Return")),0)',
         PCT_FMT),
        ("Recurring Spend / mo (manual entries only)",
         '=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[RecurringFlag],"Yes",Fact_Expenses[Date],">="&CurMonthStart,Fact_Expenses[Date],"<="&CurMonthEnd)',
         CUR_FMT),
    ]
    for i, (label, formula, fmt) in enumerate(metrics):
        rr = wr + 2 + i
        ws.cell(row=rr, column=11, value=label).font = f(9)
        c = ws.cell(row=rr, column=13, value=formula)
        c.number_format = fmt
        c.font = f(9, bold=True)

    wr2 = wr + 9
    # -- Budget Status (full category list, red = over budget) --
    ws.cell(row=wr2, column=1, value="Budget Status (red = over budget)").font = f(10, bold=True)
    for j, h in enumerate(["Category", "Status"]):
        ws.cell(row=wr2 + 1, column=1 + j, value=h).font = f(9, bold=True, color=WHITE)
        ws.cell(row=wr2 + 1, column=1 + j).fill = fill(NAVY)
    n_cat = len(EXPENSE_CATEGORIES)
    for i in range(n_cat):
        rr = wr2 + 2 + i
        ws.cell(row=rr, column=1, value=f"=Budget!A{bfr+i}").font = f(9)
        ws.cell(row=rr, column=2, value=f"=Budget!F{bfr+i}").font = f(9)
    ws.conditional_formatting.add(f"B{wr2+2}:B{wr2+1+n_cat}", CellIsRule(operator="equal", formula=['"Over Budget"'], fill=fill(RED_LIGHT), font=Font(color=RED, bold=True)))
    ws.conditional_formatting.add(f"B{wr2+2}:B{wr2+1+n_cat}", CellIsRule(operator="equal", formula=['"Near Limit"'], fill=fill(ORANGE_LIGHT), font=Font(color=ORANGE)))
    ws.conditional_formatting.add(f"B{wr2+2}:B{wr2+1+n_cat}", CellIsRule(operator="equal", formula=['"Under Budget"'], fill=fill(GREEN_LIGHT), font=Font(color=GREEN)))

    # -- Recent Transactions (last 8 rows entered) --
    ws.cell(row=wr2, column=4, value="Recent Transactions (Expenses tab only)").font = f(10, bold=True)
    for j, h in enumerate(["Date", "Vendor", "Category", "Amount"]):
        ws.cell(row=wr2 + 1, column=4 + j, value=h).font = f(9, bold=True, color=WHITE)
        ws.cell(row=wr2 + 1, column=4 + j).fill = fill(NAVY)
    for i in range(8):
        rr = wr2 + 2 + i
        idx_formula = f"COUNTA(Fact_Expenses[ExpenseID])-{i}"
        ws.cell(row=rr, column=4, value=f'=IFERROR(INDEX(Fact_Expenses[Date],{idx_formula}),"")')
        ws.cell(row=rr, column=4).number_format = DATE_FMT
        ws.cell(row=rr, column=5, value=f'=IFERROR(INDEX(Fact_Expenses[VendorName],{idx_formula}),"")')
        ws.cell(row=rr, column=6, value=f'=IFERROR(INDEX(Fact_Expenses[CategoryName],{idx_formula}),"")')
        amt = ws.cell(row=rr, column=7, value=f'=IFERROR(INDEX(Fact_Expenses[Amount],{idx_formula}),"")')
        amt.number_format = CUR_FMT
        for c in (4, 5, 6, 7):
            ws.cell(row=rr, column=c).font = f(9)

    # -- Cut-Back Analyzer -- (leave room below the taller of Budget Status / Recent Transactions)
    cba_row = wr2 + 2 + max(n_cat, 8) + 1
    cba_row = style_section_header(ws, cba_row, 1, "Cut-Back Analyzer", span=16, size=12)
    for j, h in enumerate(["Category", "Monthly Cost", "Annual Cost", "Potential Savings", "Priority", "Recommendation"]):
        ws.cell(row=cba_row, column=1 + j, value=h).font = f(9, bold=True, color=WHITE)
        ws.cell(row=cba_row, column=1 + j).fill = fill(NAVY)
    def cba_category_formula(cat):
        return (f'=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[CategoryName],"{cat}",Fact_Expenses[Date],">="&EOMONTH(TodayDate,-2)+1,Fact_Expenses[Date],"<="&EOMONTH(TodayDate,-1))'
                f'+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[CategoryName],"{cat}",BankImportRaw[TransactionType],"Expense",BankImportRaw[Date],">="&EOMONTH(TodayDate,-2)+1,BankImportRaw[Date],"<="&EOMONTH(TodayDate,-1))'
                f'-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[CategoryName],"{cat}",BankImportRaw[TransactionType],"Return",BankImportRaw[Date],">="&EOMONTH(TodayDate,-2)+1,BankImportRaw[Date],"<="&EOMONTH(TodayDate,-1))')
    cba_defs = [
        ("Dining & Restaurants", cba_category_formula("Dining & Restaurants")),
        ("Entertainment", cba_category_formula("Entertainment")),
        ("Personal Care", cba_category_formula("Personal Care")),
        ("Non-Essential Wants (all categories)",
         '=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[NeedWantFlag],"Want",Fact_Expenses[Date],">="&EOMONTH(TodayDate,-2)+1,Fact_Expenses[Date],"<="&EOMONTH(TodayDate,-1))'
         '+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[NeedWantFlag],"Want",BankImportRaw[TransactionType],"Expense",BankImportRaw[Date],">="&EOMONTH(TodayDate,-2)+1,BankImportRaw[Date],"<="&EOMONTH(TodayDate,-1))'
         '-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[NeedWantFlag],"Want",BankImportRaw[TransactionType],"Return",BankImportRaw[Date],">="&EOMONTH(TodayDate,-2)+1,BankImportRaw[Date],"<="&EOMONTH(TodayDate,-1))'),
    ]
    cba_first = cba_row + 1
    for i, (label, formula) in enumerate(cba_defs):
        rr = cba_first + i
        ws.cell(row=rr, column=1, value=label).font = f(9)
        mc = ws.cell(row=rr, column=2, value=formula)
        mc.number_format = CUR_FMT
        ws.cell(row=rr, column=3, value=f"=B{rr}*12")
        ws.cell(row=rr, column=3).number_format = CUR_FMT
        ws.cell(row=rr, column=4, value=f"=C{rr}*0.3")
        ws.cell(row=rr, column=4).number_format = CUR_FMT
        ws.cell(row=rr, column=5, value=f"=RANK(D{rr},$D${cba_first}:$D${cba_first+len(cba_defs)-1})")
        ws.cell(row=rr, column=6,
                value=f'=IF(D{rr}>500,"High-impact — consider cutting significantly",IF(D{rr}>150,"Moderate — trim where possible","Low-impact — monitor only"))')
        for c in (1, 5, 6):
            ws.cell(row=rr, column=c).font = f(9)
        ws.cell(row=rr, column=6).alignment = LEFT_WRAP

    ws.sheet_view.zoomScale = 90
    return {}
