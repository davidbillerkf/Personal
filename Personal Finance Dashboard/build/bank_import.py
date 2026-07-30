"""Bank Import sheet — paste a bank/credit-card CSV export and it's cleaned, categorized
(via your Vendor list), split into Expense/Income, and counted directly in every KPI, chart,
Budget line, and Vendor total on the Dashboard. This is a permanent ledger for bank-derived
transactions (like Expenses/Income are for manually entered ones) — rows are never copied
elsewhere or deleted. No Power Query setup required to use this."""
from openpyxl.workbook.defined_name import DefinedName
from common import *


def build_bank_import(wb):
    ws = wb.create_sheet("Bank Import")
    style_sheet_title(ws, "Bank Import", span_cols=11)
    ws.sheet_view.showGridLines = False
    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 15
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["D"].width = 22

    r = 4
    steps = [
        "HOW TO USE THIS TAB",
        "1. Export a CSV of transactions from your bank or credit card's website (usually Accounts > Download/Export).",
        "2. Open that CSV, copy its Date / Description / Amount columns, and paste them into columns A, B, C of the"
        " table below — into the first blank row, not over existing rows — paste values only (Home > Paste > Values)"
        " so formatting doesn't fight the table.",
        "3. Set 'Sign Convention' below to match how your export shows amounts (see the dropdown note).",
        "4. Column D (Vendor) tries to auto-fill itself from any earlier row in this tab with the exact same bank"
        " description — if it's blank (new or slightly different wording), pick a Vendor from the dropdown."
        " Category (E) then auto-fills from that vendor's default category; override it if needed.",
        "That's it — every row here counts immediately in the Dashboard KPIs, charts, Budget, and Vendor totals."
        " There's no copy step: this tab IS the ledger for bank-derived transactions, the same way Expenses/Income"
        " are the ledger for anything you enter by hand. Don't delete rows after importing — deleting a row removes"
        " that transaction from every report.",
        "Note: the auto-fill in step 4 only recognizes an EXACT repeat of a description string. Recurring bills and"
        " subscriptions usually post with identical text every time, so those learn fast. If your bank appends a"
        " unique date or reference number to every line, matching won't catch it — you'll just pick the vendor"
        " again, same as the first time.",
        "Note: the Dashboard's 'Top Expenses' and 'Recent Transactions' widgets only look at the Expenses tab"
        " (a limitation of ranking individual transactions across two separate tables) — everything else (KPIs,"
        " Budget, Vendor totals, charts, Cut-Back Analyzer) includes this tab.",
    ]
    for i, s in enumerate(steps):
        rr = r + i
        c = ws.cell(row=rr, column=1, value=s)
        c.font = f(11, bold=True) if i == 0 else f(10)
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=11)
        c.alignment = LEFT_WRAP
        ws.row_dimensions[rr].height = 30 if i > 0 else 20

    toggle_row = r + len(steps) + 1
    ws.cell(row=toggle_row, column=1, value="Sign Convention for this import:").font = f(10, bold=True)
    ws.merge_cells(start_row=toggle_row, start_column=1, end_row=toggle_row, end_column=3)
    toggle_cell = ws.cell(row=toggle_row, column=4, value="Negative = Expense")
    toggle_cell.font = f(10, bold=True, color=BLUE_ACCENT)
    toggle_cell.fill = fill(LIGHT_GRAY)
    dv = DataValidation(type="list", formula1='"Negative = Expense,Positive = Expense"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(toggle_cell.coordinate)
    ws.cell(row=toggle_row, column=5,
            value="Most bank/debit exports show purchases as negative numbers (“Negative = Expense”)."
                  " Most credit-card exports show charges as positive numbers and payments as negative"
                  " (“Positive = Expense”). Check one real row against your bank statement to confirm.")
    ws.cell(row=toggle_row, column=5).font = f(9, italic=True, color="6B7280")
    ws.merge_cells(start_row=toggle_row, start_column=5, end_row=toggle_row, end_column=11)
    ws.cell(row=toggle_row, column=5).alignment = LEFT_WRAP
    ws.row_dimensions[toggle_row].height = 40
    wb.defined_names["SignConvention"] = DefinedName("SignConvention", attr_text=f"'Bank Import'!$D${toggle_row}")

    r2 = toggle_row + 2
    r2 = style_section_header(ws, r2, 1,
        "Paste your CSV's Date / Description / Amount into columns A-C — Vendor (D) auto-fills from earlier rows when it recognizes the description",
        span=11)
    headers = ["Date", "Description", "Amount", "VendorName", "CategoryName", "AccountName",
               "PaymentMethod", "TransactionType", "AbsAmount", "NeedWantFlag", "Notes"]
    n_rows = 500
    fr = r2 + 1
    for j, h in enumerate(headers):
        c = ws.cell(row=r2, column=1 + j, value=h)
        c.font = f(10, bold=True, color=WHITE)
        c.fill = fill(NAVY)
        c.alignment = CENTER_WRAP
    lr = fr + n_rows - 1
    ref = f"A{r2}:K{lr}"
    tab = Table(displayName="BankImportRaw", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)

    for rr in range(fr, lr + 1):
        ws.cell(row=rr, column=1).number_format = DATE_FMT
        ws.cell(row=rr, column=3).number_format = CUR_FMT
        # Vendor "remembers" prior imports: exact-match this row's raw Description against Description
        # text in EARLIER rows of this same table only (never itself or later rows — that would create a
        # circular reference, since this column feeds Fact-level totals that this same formula could see).
        # Plain scalar INDEX/MATCH over a growing range — no arrays, safe in any Excel version.
        if rr == fr:
            ws.cell(row=rr, column=4, value='=""')
        else:
            ws.cell(row=rr, column=4,
                    value=(f'=IF($B{rr}="","",'
                           f'IFERROR(INDEX($D${fr}:$D${rr-1},MATCH($B{rr},$B${fr}:$B${rr-1},0)),""))'))
        # Category auto-suggested from the picked Vendor's default category (plain, non-array INDEX/MATCH)
        ws.cell(row=rr, column=5,
                value=f'=IFERROR(INDEX(Dim_Vendor[DefaultCategory],MATCH($D{rr},Dim_Vendor[VendorName],0)),"")')
        ws.cell(row=rr, column=8,
                value=(f'=IF($C{rr}="","",'
                       f'IF(SignConvention="Negative = Expense",IF($C{rr}<0,"Expense","Income"),'
                       f'IF($C{rr}>0,"Expense","Income")))'))
        ws.cell(row=rr, column=9, value=f'=IF($C{rr}="","",ABS($C{rr}))')
        ws.cell(row=rr, column=9).number_format = CUR_FMT

    add_list_validation(ws, f"D{fr}:D{lr}", "=VendorNameList")
    add_list_validation(ws, f"E{fr}:E{lr}", "=ExpenseCategoryList")
    add_list_validation(ws, f"F{fr}:F{lr}", "=AccountNameList")
    add_list_validation(ws, f"G{fr}:G{lr}", "=PaymentMethodList")
    add_list_validation(ws, f"J{fr}:J{lr}", "=List_NeedWant")

    ws.conditional_formatting.add(
        f"H{fr}:H{lr}", CellIsRule(operator="equal", formula=['"Expense"'], font=Font(color=RED)))
    ws.conditional_formatting.add(
        f"H{fr}:H{lr}", CellIsRule(operator="equal", formula=['"Income"'], font=Font(color=GREEN)))
    ws.conditional_formatting.add(
        f"D{fr}:D{lr}", FormulaRule(formula=[f'AND($C{fr}<>"",$D{fr}="")'], fill=fill(ORANGE_LIGHT)))

    ws.freeze_panes = f"A{fr}"
    ws.sheet_properties.tabColor = ORANGE
    return {"fr": fr, "lr": lr, "toggle_row": toggle_row}
