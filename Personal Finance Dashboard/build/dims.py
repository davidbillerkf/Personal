import datetime as dt
from openpyxl.workbook.defined_name import DefinedName
from common import *
import seed_data as sd


def build_settings(wb):
    ws = wb.create_sheet("Settings")
    style_sheet_title(ws, "Settings", span_cols=8)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 16

    r = style_section_header(ws, 4, 1, "System Date Anchors (used by all SUMIFS date filters)", span=5)
    labels = [
        ("Today's Date", "=TODAY()", "TodayDate"),
        ("Current Month Start", "=EOMONTH(TodayDate,-1)+1", "CurMonthStart"),
        ("Current Month End", "=EOMONTH(TodayDate,0)", "CurMonthEnd"),
        ("Last Month Start", "=EOMONTH(TodayDate,-2)+1", "LastMonthStart"),
        ("Last Month End", "=EOMONTH(TodayDate,-1)", "LastMonthEnd"),
        ("Current Year Start", "=DATE(YEAR(TodayDate),1,1)", "CurYearStart"),
        ("Current Year End", "=DATE(YEAR(TodayDate),12,31)", "CurYearEnd"),
    ]
    named_cells = {}
    for i, (label, formula, name) in enumerate(labels):
        row = r + i
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = f(10)
        lc.alignment = LEFT
        vc = ws.cell(row=row, column=2, value=formula)
        vc.font = f(10, bold=True, color=BLUE_ACCENT)
        vc.number_format = DATE_FMT
        named_cells[name] = f"Settings!$B${row}"
    r2 = r + len(labels) + 1

    r2 = style_section_header(ws, r2 + 1, 1, "Dim_PaymentMethod", span=5)
    pm_headers = ["PaymentMethodID", "MethodName"]
    pm_rows = [(1, "Credit Card"), (2, "Debit Card"), (3, "Cash"), (4, "ACH"), (5, "Check")]
    pm_info = add_table(ws, r2, 1, pm_headers, "Dim_PaymentMethod", pm_rows, col_widths=[16, 20])
    r3 = pm_info["end_row"] + 2

    r3 = style_section_header(ws, r3 + 1, 1, "Controlled Vocabulary Lists (used for data-validation dropdowns)", span=5)
    lists = {
        "List_YesNo": ["Yes", "No"],
        "List_NeedWant": ["Need", "Want"],
    }
    col = 1
    start_row = r3
    list_ranges = {}
    for name, values in lists.items():
        ws.cell(row=start_row, column=col, value=name).font = f(10, bold=True, color=WHITE)
        ws.cell(row=start_row, column=col).fill = fill(NAVY)
        for i, v in enumerate(values):
            ws.cell(row=start_row + 1 + i, column=col, value=v).font = f(10)
        first = start_row + 1
        last = start_row + len(values)
        list_ranges[name] = f"Settings!${get_column_letter(col)}${first}:${get_column_letter(col)}${last}"
        col += 1

    for name, ref in named_cells.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    for name, ref in list_ranges.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    wb.defined_names["PaymentMethodList"] = DefinedName(
        "PaymentMethodList", attr_text="Dim_PaymentMethod[MethodName]")

    ws.sheet_properties.tabColor = "6B7280"
    return {"named": named_cells, "lists": list_ranges, "pm_info": pm_info}


def build_dim_date(wb):
    """Hidden helper sheet: Dim_Date, DateID = INT(date serial). 2025-2027."""
    ws = wb.create_sheet("Dim_Date")
    headers = ["DateID", "Date", "Month", "MonthNumber", "Quarter", "Year"]
    for j, h in enumerate(headers):
        c = ws.cell(row=1, column=1 + j, value=h)
        c.font = f(10, bold=True, color=WHITE)
        c.fill = fill(NAVY)
    start = dt.date(2025, 1, 1)
    end = dt.date(2027, 12, 31)
    n_days = (end - start).days + 1
    row = 2
    for i in range(n_days):
        this_date = start + dt.timedelta(days=i)
        r = row + i
        ws.cell(row=r, column=1, value=f"=INT(B{r})")
        dc = ws.cell(row=r, column=2, value=this_date)
        dc.number_format = DATE_FMT
        ws.cell(row=r, column=3, value=f'=TEXT(B{r},"mmm")')
        ws.cell(row=r, column=4, value=f"=MONTH(B{r})")
        ws.cell(row=r, column=5, value=f"=ROUNDUP(MONTH(B{r})/3,0)")
        ws.cell(row=r, column=6, value=f"=YEAR(B{r})")
    end_row = row + n_days - 1
    ref = f"A1:F{end_row}"
    tab = Table(displayName="Dim_Date", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)
    ws.sheet_state = "hidden"
    return {"end_row": end_row}


def build_categories(wb):
    ws = wb.create_sheet("Categories")
    style_sheet_title(ws, "Categories", span_cols=6)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16

    r = style_section_header(ws, 4, 1, "Dim_Category  —  add a row here to make a new category available everywhere", span=4)
    cat_info = add_table(ws, r, 1, ["CategoryID", "CategoryName", "CategoryType", "EssentialCategory"],
                          "Dim_Category", sd.CATEGORIES, col_widths=[12, 24, 14, 16])

    wb.defined_names["CategoryNameList"] = DefinedName(
        "CategoryNameList", attr_text="Dim_Category[CategoryName]")
    wb.defined_names["ExpenseCategoryList"] = DefinedName(
        "ExpenseCategoryList", attr_text="Dim_Category[CategoryName]")
    ws.sheet_properties.tabColor = "6B7280"
    return {"cat_info": cat_info}


def build_vendors(wb):
    ws = wb.create_sheet("Vendors")
    style_sheet_title(ws, "Vendors", span_cols=6)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 13

    r = style_section_header(ws, 4, 1,
        "Dim_Vendor  —  add a row to make a new vendor available in Expenses; set/edit DefaultCategory (dropdown) so Bank Import can auto-suggest a category",
        span=6)
    ven_info = add_table(ws, r, 1, ["VendorID", "VendorName", "DefaultCategory", "DefaultCategoryID", "CurMonthSpend", "YTDSpend"],
                          "Dim_Vendor", [[vid, name, cat, None, None, None] for (vid, name, cat) in sd.VENDORS],
                          col_widths=[12, 26, 20, 15, 15, 13])
    fr, lr = ven_info["first_data_row"], ven_info["end_row"]
    for rr in range(fr, lr + 1):
        ws.cell(row=rr, column=4,
                value=f'=IFERROR(INDEX(Dim_Category[CategoryID],MATCH($C{rr},Dim_Category[CategoryName],0)),"")')
        ws.cell(row=rr, column=5,
                value=(f'=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[VendorName],$B{rr},Fact_Expenses[Date],">="&CurMonthStart,Fact_Expenses[Date],"<="&CurMonthEnd)'
                       f'+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[VendorName],$B{rr},BankImportRaw[TransactionType],"Expense",BankImportRaw[Date],">="&CurMonthStart,BankImportRaw[Date],"<="&CurMonthEnd)'
                       f'-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[VendorName],$B{rr},BankImportRaw[TransactionType],"Return",BankImportRaw[Date],">="&CurMonthStart,BankImportRaw[Date],"<="&CurMonthEnd)'))
        ws.cell(row=rr, column=5).number_format = CUR_FMT
        ws.cell(row=rr, column=6,
                value=(f'=SUMIFS(Fact_Expenses[Amount],Fact_Expenses[VendorName],$B{rr},Fact_Expenses[Date],">="&CurYearStart,Fact_Expenses[Date],"<="&CurYearEnd)'
                       f'+SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[VendorName],$B{rr},BankImportRaw[TransactionType],"Expense",BankImportRaw[Date],">="&CurYearStart,BankImportRaw[Date],"<="&CurYearEnd)'
                       f'-SUMIFS(BankImportRaw[AbsAmount],BankImportRaw[VendorName],$B{rr},BankImportRaw[TransactionType],"Return",BankImportRaw[Date],">="&CurYearStart,BankImportRaw[Date],"<="&CurYearEnd)'))
        ws.cell(row=rr, column=6).number_format = CUR_FMT
    add_list_validation(ws, f"C{fr}:C{lr+50}", "=CategoryNameList")
    ws.column_dimensions["D"].hidden = True
    wb.defined_names["VendorNameList"] = DefinedName(
        "VendorNameList", attr_text="Dim_Vendor[VendorName]")
    ws.sheet_properties.tabColor = "6B7280"
    return {"ven_info": ven_info}


def build_accounts(wb):
    ws = wb.create_sheet("Accounts")
    style_sheet_title(ws, "Accounts", span_cols=6)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16

    r = style_section_header(ws, 4, 1, "Dim_Account  —  add a row here to make a new account available everywhere", span=4)
    acc_info = add_table(ws, r, 1, ["AccountID", "AccountName", "AccountType", "Balance"],
                          "Dim_Account", sd.ACCOUNTS, col_widths=[12, 20, 16, 16],
                          number_formats={3: CUR_FMT})
    for rr in range(acc_info["first_data_row"], acc_info["end_row"] + 1):
        ws.cell(row=rr, column=4).number_format = CUR_FMT
        ws.cell(row=rr, column=4).font = f(10, color=BLUE_ACCENT)

    r2 = acc_info["end_row"] + 3
    r2 = style_section_header(ws, r2, 1, "Account Totals", span=4)
    ws.cell(row=r2, column=1, value="Total All Accounts").font = f(10, bold=True)
    ws.cell(row=r2, column=2, value="=SUM(Dim_Account[Balance])").font = f(10, bold=True)
    ws.cell(row=r2, column=2).number_format = CUR_FMT

    wb.defined_names["AccountNameList"] = DefinedName(
        "AccountNameList", attr_text="Dim_Account[AccountName]")
    ws.sheet_properties.tabColor = "6B7280"
    return {"acc_info": acc_info}
