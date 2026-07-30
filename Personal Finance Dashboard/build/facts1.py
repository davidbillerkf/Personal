"""Income and Expenses sheets (simplified planner scope)."""
from openpyxl.workbook.defined_name import DefinedName
from common import *
import seed_data as sd


def build_income(wb):
    ws = wb.create_sheet("Income")
    style_sheet_title(ws, "Income", span_cols=14)
    ws.sheet_view.showGridLines = False

    r = style_section_header(ws, 4, 1, "Dim_Source", span=2)
    sources = [(1, "Employer Payroll Inc"), (2, "Other")]
    src_info = add_table(ws, r, 1, ["SourceID", "SourceName"], "Dim_Source", sources, col_widths=[10, 24])
    wb.defined_names["SourceNameList"] = DefinedName(
        "SourceNameList", attr_text="Dim_Source[SourceName]")

    r2 = src_info["end_row"] + 3
    r2 = style_section_header(ws, r2, 1,
        "Fact_Income  —  enter each paycheck/deposit here. Dropdowns keep entries consistent; ID columns (hidden) auto-lookup for reporting.",
        span=14)
    headers = ["IncomeID", "Date", "SourceName", "SourceID", "CategoryName", "CategoryID",
               "AccountName", "AccountID", "GrossAmount", "TaxesWithheld", "NetAmount", "RecurringFlag", "Notes",
               "Description"]
    data_rows = []
    for row in sd.INCOME_ROWS:
        date_, source, cat, acct, gross, tax, rec, notes = row
        data_rows.append([None, date_, source, None, cat, None, acct, None, gross, tax, None, rec, notes, None])
    inc_info = add_table(ws, r2, 1, headers, "Fact_Income", data_rows,
                          col_widths=[10, 12, 22, 10, 16, 10, 20, 10, 13, 13, 13, 13, 24, 30])

    fr, lr = inc_info["first_data_row"], inc_info["end_row"]
    for rr in range(fr, lr + 1):
        ws.cell(row=rr, column=1, value=f"=ROW()-{fr-1}")
        ws.cell(row=rr, column=2).number_format = DATE_FMT
        ws.cell(row=rr, column=4, value=f'=IFERROR(INDEX(Dim_Source[SourceID],MATCH($C{rr},Dim_Source[SourceName],0)),"")')
        ws.cell(row=rr, column=6, value=f'=IFERROR(INDEX(Dim_Category[CategoryID],MATCH($E{rr},Dim_Category[CategoryName],0)),"")')
        ws.cell(row=rr, column=8, value=f'=IFERROR(INDEX(Dim_Account[AccountID],MATCH($G{rr},Dim_Account[AccountName],0)),"")')
        for col in (9, 10, 11):
            ws.cell(row=rr, column=col).number_format = CUR_FMT
        ws.cell(row=rr, column=11, value=f"=$I{rr}-$J{rr}")
    for col_letter in ("D", "F", "H"):
        ws.column_dimensions[col_letter].hidden = True

    buf = 300
    add_list_validation(ws, f"C{fr}:C{lr+buf}", "=SourceNameList")
    add_list_validation(ws, f"E{fr}:E{lr+buf}", "=ExpenseCategoryList")
    add_list_validation(ws, f"G{fr}:G{lr+buf}", "=AccountNameList")
    add_list_validation(ws, f"L{fr}:L{lr+buf}", "=List_YesNo")

    ws.freeze_panes = f"A{fr}"
    ws.sheet_properties.tabColor = GREEN
    return {"inc_info": inc_info}


def build_expenses(wb):
    ws = wb.create_sheet("Expenses")
    style_sheet_title(ws, "Expenses", span_cols=16)
    ws.sheet_view.showGridLines = False

    r = style_section_header(ws, 4, 1,
        "Fact_Expenses  —  every purchase goes here. Pick Vendor/Category/Account/Payment Method from dropdowns.",
        span=16)
    headers = ["ExpenseID", "Date", "VendorName", "VendorID", "CategoryName", "CategoryID",
               "AccountName", "AccountID", "PaymentMethod", "PaymentMethodID",
               "Amount", "NeedWantFlag", "RecurringFlag", "Notes", "Description", "CurMonthAmt"]
    data_rows = []
    for row in sd.EXPENSE_ROWS:
        (date_, vendor, cat, acct, pm, amount, nw, rec, notes) = row
        data_rows.append([None, date_, vendor, None, cat, None, acct, None, pm, None,
                           amount, nw, rec, notes, None, None])
    exp_info = add_table(ws, r, 1, headers, "Fact_Expenses", data_rows,
                          col_widths=[10, 12, 20, 10, 20, 10, 16, 10, 14, 14, 12, 11, 11, 22, 30, 13])

    fr, lr = exp_info["first_data_row"], exp_info["end_row"]
    for rr in range(fr, lr + 1):
        ws.cell(row=rr, column=1, value=f"=ROW()-{fr-1}")
        ws.cell(row=rr, column=2).number_format = DATE_FMT
        ws.cell(row=rr, column=4, value=f'=IFERROR(INDEX(Dim_Vendor[VendorID],MATCH($C{rr},Dim_Vendor[VendorName],0)),"")')
        ws.cell(row=rr, column=6, value=f'=IFERROR(INDEX(Dim_Category[CategoryID],MATCH($E{rr},Dim_Category[CategoryName],0)),"")')
        ws.cell(row=rr, column=8, value=f'=IFERROR(INDEX(Dim_Account[AccountID],MATCH($G{rr},Dim_Account[AccountName],0)),"")')
        ws.cell(row=rr, column=10, value=f'=IFERROR(INDEX(Dim_PaymentMethod[PaymentMethodID],MATCH($I{rr},Dim_PaymentMethod[MethodName],0)),"")')
        ws.cell(row=rr, column=11).number_format = CUR_FMT
        ws.cell(row=rr, column=16,
                value=f'=IF(AND($B{rr}>=CurMonthStart,$B{rr}<=CurMonthEnd),$K{rr},0)')
        ws.cell(row=rr, column=16).number_format = CUR_FMT
    for col_letter in ("D", "F", "H", "J", "P"):
        ws.column_dimensions[col_letter].hidden = True

    buf = 300
    add_list_validation(ws, f"C{fr}:C{lr+buf}", "=VendorNameList")
    add_list_validation(ws, f"E{fr}:E{lr+buf}", "=ExpenseCategoryList")
    add_list_validation(ws, f"G{fr}:G{lr+buf}", "=AccountNameList")
    add_list_validation(ws, f"I{fr}:I{lr+buf}", "=PaymentMethodList")
    add_list_validation(ws, f"L{fr}:L{lr+buf}", "=List_NeedWant")
    add_list_validation(ws, f"M{fr}:M{lr+buf}", "=List_YesNo")

    ws.conditional_formatting.add(
        f"K{fr}:K{lr}",
        CellIsRule(operator="greaterThan", formula=["200"], fill=fill(ORANGE_LIGHT)))

    ws.freeze_panes = f"A{fr}"
    ws.sheet_properties.tabColor = RED
    return {"exp_info": exp_info}
